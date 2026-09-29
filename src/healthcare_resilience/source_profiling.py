"""Structural profiling for contracted CSV tables and XLSX worksheets."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


class SourceTableProfilingError(ValueError):
    """Raised when a source cannot be profiled without violating its contract."""


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise SourceTableProfilingError(message)


def ordered_header_sha256(header: list[str]) -> str:
    """Hash an ordered header with the algorithm named by the public contract."""

    payload = json.dumps(
        header,
        ensure_ascii=False,
        allow_nan=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _row_fingerprint(row: Iterable[Any], *, text_only: bool) -> bytes:
    if text_only:
        values = list(row)
    else:
        values = []
        for value in row:
            if isinstance(value, (datetime, date)):
                values.append({"datetime": value.isoformat()})
            elif isinstance(value, float):
                _require(math.isfinite(value), "source row contains a non-finite number")
                values.append(value)
            else:
                values.append(value)
    payload = json.dumps(
        values,
        ensure_ascii=False,
        allow_nan=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).digest()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _key_value(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _parse_datetime_text(value: str) -> None:
    candidates = (
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d",
        "%Y/%m/%d %H:%M:%S",
        "%m/%d/%Y",
        "%m/%d/%Y %H:%M:%S",
    )
    try:
        datetime.fromisoformat(value)
        return
    except ValueError:
        pass
    for candidate in candidates:
        try:
            datetime.strptime(value, candidate)
            return
        except ValueError:
            continue
    raise ValueError("unsupported datetime text")


def _validate_parser_value(value: Any, parser_type: str) -> None:
    if parser_type == "STRING":
        _require(isinstance(value, str), "value is not text")
        return
    if parser_type == "INTEGER":
        if isinstance(value, bool):
            raise ValueError("boolean is not an integer")
        if isinstance(value, int):
            return
        if isinstance(value, float) and math.isfinite(value) and value.is_integer():
            return
        if isinstance(value, str):
            int(value.strip())
            return
        raise ValueError("value is not an integer")
    if parser_type == "NUMBER":
        if isinstance(value, bool):
            raise ValueError("boolean is not a number")
        if isinstance(value, (int, float)):
            if not math.isfinite(float(value)):
                raise ValueError("number is not finite")
            return
        if isinstance(value, str):
            try:
                parsed = Decimal(value.strip())
            except InvalidOperation as error:
                raise ValueError("value is not numeric") from error
            if not parsed.is_finite():
                raise ValueError("number is not finite")
            return
        raise ValueError("value is not numeric")
    if parser_type == "DATETIME":
        if isinstance(value, (datetime, date)):
            return
        if isinstance(value, str):
            _parse_datetime_text(value.strip())
            return
        raise ValueError("value is not a datetime")
    raise ValueError(f"unsupported parser type {parser_type}")


class _TableAccumulator:
    def __init__(
        self, contract: dict[str, Any], header: list[str], *, text_rows: bool
    ) -> None:
        self.contract = contract
        self.header = header
        self.text_rows = text_rows
        self.row_count = 0
        self.row_widths: set[int] = set()
        self.row_fingerprints: set[bytes] = set()
        self.exact_duplicate_count = 0
        self.key_nonblank_count = 0
        self.key_values: set[tuple[str, ...]] = set()

        named = [name for name in header if name]
        _require(
            len(named) == len(set(named)),
            f"{contract['id']}: named header contains duplicates",
        )
        self.column_indexes = {name: index for index, name in enumerate(header) if name}
        required_names = [item["name"] for item in contract["required_columns"]]
        missing = sorted(set(required_names) - set(self.column_indexes))
        _require(not missing, f"{contract['id']}: missing required columns {missing}")
        self.required_columns = [
            (self.column_indexes[item["name"]], item)
            for item in contract["required_columns"]
        ]
        self.key_indexes = [
            self.column_indexes[name] for name in contract["record_key"]["columns"]
        ]

    def add_row(self, row: tuple[Any, ...] | list[Any], row_number: int) -> None:
        self.row_count += 1
        self.row_widths.add(len(row))
        fingerprint = _row_fingerprint(row, text_only=self.text_rows)
        if fingerprint in self.row_fingerprints:
            self.exact_duplicate_count += 1
        else:
            self.row_fingerprints.add(fingerprint)

        for index, column in self.required_columns:
            value = row[index] if index < len(row) else None
            if _is_blank(value):
                _require(
                    column["null_allowed"],
                    f"{self.contract['id']}: {column['name']} is blank at row {row_number}",
                )
                continue
            try:
                _validate_parser_value(value, column["parser_type"])
            except (TypeError, ValueError) as error:
                raise SourceTableProfilingError(
                    f"{self.contract['id']}: {column['name']} cannot be parsed as "
                    f"{column['parser_type']} at row {row_number}"
                ) from error

        key_values = [row[index] if index < len(row) else None for index in self.key_indexes]
        if all(not _is_blank(value) for value in key_values):
            self.key_nonblank_count += 1
            self.key_values.add(tuple(_key_value(value) for value in key_values))

    def finish(self) -> dict[str, Any]:
        _require(self.row_widths, f"{self.contract['id']}: table has no data rows")
        _require(
            len(self.row_widths) == 1,
            f"{self.contract['id']}: inconsistent data-row widths {sorted(self.row_widths)}",
        )
        data_row_field_count = next(iter(self.row_widths))
        key_distinct_count = len(self.key_values)
        return {
            "id": self.contract["id"],
            "source_id": self.contract["source_id"],
            "filename": self.contract["filename"],
            "container": self.contract["container"],
            "sheet_name": self.contract["sheet_name"],
            "encoding": self.contract["encoding"],
            "byte_order_mark": self.contract["byte_order_mark"],
            "csv_dialect": self.contract["csv_dialect"],
            "header_row": self.contract["header_row"],
            "data_row_count": self.row_count,
            "header_cell_count": len(self.header),
            "named_column_count": sum(bool(name) for name in self.header),
            "data_row_field_count": data_row_field_count,
            "unnamed_header_positions": [
                index for index, name in enumerate(self.header, start=1) if not name
            ],
            "header_sha256": ordered_header_sha256(self.header),
            "required_column_types": {
                item["name"]: item["parser_type"]
                for item in self.contract["required_columns"]
            },
            "record_key_evidence": {
                "columns": self.contract["record_key"]["columns"],
                "nonblank_rows": self.key_nonblank_count,
                "distinct_count": key_distinct_count,
                "duplicate_rows_beyond_first": (
                    self.key_nonblank_count - key_distinct_count
                ),
            },
            "exact_duplicate_rows_beyond_first": self.exact_duplicate_count,
        }


def _line_ending(raw: bytes, table_id: str) -> str:
    crlf_count = raw.count(b"\r\n")
    bare_lf_count = raw.count(b"\n") - crlf_count
    bare_cr_count = raw.count(b"\r") - crlf_count
    kinds = sum(count > 0 for count in (crlf_count, bare_lf_count, bare_cr_count))
    _require(kinds == 1, f"{table_id}: missing or mixed CSV line endings")
    _require(bare_cr_count == 0, f"{table_id}: unsupported CR-only line endings")
    return "CRLF" if crlf_count else "LF"


def _profile_csv(path: Path, contract: dict[str, Any]) -> dict[str, Any]:
    raw = path.read_bytes()
    bom_present = raw.startswith(b"\xef\xbb\xbf")
    observed_bom = "PRESENT" if bom_present else "ABSENT"
    encoding = contract["encoding"]
    codec = {
        "UTF-8": "utf-8-sig" if bom_present else "utf-8",
        "WINDOWS-1252": "cp1252",
    }[encoding]
    try:
        text = raw.decode(codec, errors="strict")
    except UnicodeDecodeError as error:
        raise SourceTableProfilingError(
            f"{contract['id']}: file is not valid {encoding}"
        ) from error

    dialect = contract["csv_dialect"]
    reader = csv.reader(
        io.StringIO(text, newline=""),
        delimiter=dialect["delimiter"],
        quotechar=dialect["quote_character"],
        strict=True,
    )
    try:
        header = next(reader)
    except (StopIteration, csv.Error) as error:
        raise SourceTableProfilingError(
            f"{contract['id']}: CSV header cannot be read"
        ) from error
    accumulator = _TableAccumulator(contract, header, text_rows=True)
    try:
        for row_number, row in enumerate(reader, start=2):
            accumulator.add_row(row, row_number)
    except csv.Error as error:
        raise SourceTableProfilingError(
            f"{contract['id']}: malformed CSV data"
        ) from error
    result = accumulator.finish()
    result["byte_order_mark"] = observed_bom
    result["csv_dialect"] = {**dialect, "line_ending": _line_ending(raw, contract["id"])}
    return result


def _profile_worksheet(worksheet: Any, contract: dict[str, Any]) -> dict[str, Any]:
    _require(contract["header_row"] == 1, f"{contract['id']}: only header row 1 is supported")
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_header = next(rows)
    except StopIteration as error:
        raise SourceTableProfilingError(
            f"{contract['id']}: worksheet header cannot be read"
        ) from error
    header = ["" if value is None else str(value) for value in raw_header]
    accumulator = _TableAccumulator(contract, header, text_rows=False)
    for row_number, row in enumerate(rows, start=2):
        accumulator.add_row(row, row_number)
    return accumulator.finish()


def profile_source_tables(
    raw_root: Path,
    table_contract: dict[str, Any],
    source_contract_sha256: str,
    source_table_contract_sha256: str,
) -> dict[str, Any]:
    """Profile contracted source structure without returning source values."""

    _require(raw_root.is_dir(), "snapshot raw directory is missing")
    tables = table_contract["table_contracts"]
    workbook_contracts: dict[str, list[dict[str, Any]]] = {}
    for contract in tables:
        path = raw_root / contract["filename"]
        _require(path.is_file(), f"{contract['id']}: source file is missing")
        _require(not path.is_symlink(), f"{contract['id']}: symbolic links are not allowed")
        if contract["container"] == "XLSX":
            workbook_contracts.setdefault(contract["filename"], []).append(contract)

    workbook_profiles: dict[str, dict[str, dict[str, Any]]] = {}
    for filename, contracts in workbook_contracts.items():
        workbook = load_workbook(raw_root / filename, read_only=True, data_only=True)
        try:
            expected_sheets = [contract["sheet_name"] for contract in contracts]
            _require(
                workbook.sheetnames == expected_sheets,
                f"{filename}: missing, unexpected, or reordered worksheet",
            )
            workbook_profiles[filename] = {
                contract["sheet_name"]: _profile_worksheet(
                    workbook[contract["sheet_name"]], contract
                )
                for contract in contracts
            }
        finally:
            workbook.close()

    profiles = []
    for contract in tables:
        if contract["container"] == "CSV":
            profiles.append(_profile_csv(raw_root / contract["filename"], contract))
        else:
            profiles.append(
                workbook_profiles[contract["filename"]][contract["sheet_name"]]
            )
    return {
        "schema_version": "1.0.0",
        "profile_type": "SOURCE_TABLE_STRUCTURE",
        "status": "VALID",
        "source_snapshot_id": table_contract["snapshot_id"],
        "source_contract_sha256": source_contract_sha256,
        "source_table_contract_sha256": source_table_contract_sha256,
        "table_count": len(profiles),
        "tables": profiles,
    }


def write_source_table_profile(path: Path, profile: dict[str, Any]) -> None:
    """Atomically write a deterministic profile outside the source snapshot."""

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.name}.tmp")
    text = json.dumps(profile, ensure_ascii=True, allow_nan=False, indent=2)
    temporary_path.write_text(f"{text}\n", encoding="utf-8", newline="\n")
    temporary_path.replace(path)
