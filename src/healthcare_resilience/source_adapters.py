"""Contract-driven CSV and XLSX readers for source-preserving staging."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .source_profiling import ordered_header_sha256
from .staging import extract_staging_rows


class SourceAdapterError(ValueError):
    """Raised when a source file cannot satisfy its table contract."""


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise SourceAdapterError(message)


def _line_ending(raw: bytes, table_id: str) -> str:
    crlf_count = raw.count(b"\r\n")
    bare_lf_count = raw.count(b"\n") - crlf_count
    bare_cr_count = raw.count(b"\r") - crlf_count
    kinds = sum(count > 0 for count in (crlf_count, bare_lf_count, bare_cr_count))
    _require(kinds == 1, f"{table_id}: missing or mixed CSV line endings")
    _require(bare_cr_count == 0, f"{table_id}: unsupported CR-only line endings")
    return "CRLF" if crlf_count else "LF"


def _validate_header(header: list[str], contract: dict[str, Any]) -> None:
    table_id = contract["id"]
    named = [name for name in header if name]
    unnamed_positions = [
        index for index, name in enumerate(header, start=1) if not name
    ]
    _require(
        len(header) == contract["header_cell_count"],
        f"{table_id}: header cell count differs from contract",
    )
    _require(
        len(named) == contract["named_column_count"],
        f"{table_id}: named column count differs from contract",
    )
    _require(
        len(named) == len(set(named)),
        f"{table_id}: named header contains duplicates",
    )
    _require(
        unnamed_positions == contract["unnamed_header_positions"],
        f"{table_id}: unnamed header positions differ from contract",
    )
    _require(
        ordered_header_sha256(header) == contract["header_sha256"],
        f"{table_id}: ordered header hash differs from contract",
    )


def _rows_to_mappings(
    rows: Iterable[Sequence[Any]], header: list[str], contract: dict[str, Any]
) -> list[dict[str, Any]]:
    table_id = contract["id"]
    column_indexes = {name: index for index, name in enumerate(header) if name}
    retained_names = [column["name"] for column in contract["required_columns"]]
    missing = sorted(set(retained_names) - set(column_indexes))
    _require(not missing, f"{table_id}: missing required columns {missing}")
    retained_indexes = [(name, column_indexes[name]) for name in retained_names]
    mappings: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        _require(
            len(row) == contract["data_row_field_count"],
            f"{table_id}: row {row_number} width differs from contract",
        )
        mapping = {name: row[index] for name, index in retained_indexes}
        mappings.append(mapping)
    _require(
        len(mappings) == contract["data_row_count"],
        f"{table_id}: data row count differs from contract",
    )
    return mappings


def _read_csv_rows(path: Path, contract: dict[str, Any]) -> list[dict[str, Any]]:
    raw = path.read_bytes()
    bom_present = raw.startswith(b"\xef\xbb\xbf")
    observed_bom = "PRESENT" if bom_present else "ABSENT"
    _require(
        observed_bom == contract["byte_order_mark"],
        f"{contract['id']}: byte-order mark differs from contract",
    )

    encoding = contract["encoding"]
    codec = {
        "UTF-8": "utf-8-sig" if bom_present else "utf-8",
        "WINDOWS-1252": "cp1252",
    }.get(encoding)
    _require(codec is not None, f"{contract['id']}: unsupported CSV encoding")
    try:
        text = raw.decode(codec, errors="strict")
    except UnicodeDecodeError as error:
        raise SourceAdapterError(
            f"{contract['id']}: file is not valid {encoding}"
        ) from error

    dialect = contract["csv_dialect"]
    _require(
        _line_ending(raw, contract["id"]) == dialect["line_ending"],
        f"{contract['id']}: CSV line ending differs from contract",
    )
    reader = csv.reader(
        io.StringIO(text, newline=""),
        delimiter=dialect["delimiter"],
        quotechar=dialect["quote_character"],
        strict=True,
    )
    try:
        header = next(reader)
    except (StopIteration, csv.Error) as error:
        raise SourceAdapterError(
            f"{contract['id']}: CSV cannot be read"
        ) from error
    _validate_header(header, contract)
    try:
        return _rows_to_mappings(reader, header, contract)
    except csv.Error as error:
        raise SourceAdapterError(
            f"{contract['id']}: CSV cannot be read"
        ) from error


def _read_xlsx_rows(path: Path, contract: dict[str, Any]) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = contract["sheet_name"]
        _require(
            sheet_name in workbook.sheetnames,
            f"{contract['id']}: contracted worksheet is missing",
        )
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_header = next(rows)
        except StopIteration as error:
            raise SourceAdapterError(
                f"{contract['id']}: worksheet header cannot be read"
            ) from error
        header = ["" if value is None else str(value) for value in raw_header]
        _validate_header(header, contract)
        return _rows_to_mappings(rows, header, contract)
    finally:
        workbook.close()


def read_source_rows(
    path: Path, source_table_contract: dict[str, Any]
) -> list[dict[str, Any]]:
    """Read retained columns without filtering or changing source order.

    Every physical row is still checked against the complete contracted width.
    Columns outside ``required_columns`` are validated structurally but are not
    retained in memory.
    """

    table_id = source_table_contract["id"]
    _require(path.is_file(), f"{table_id}: source file is missing")
    _require(not path.is_symlink(), f"{table_id}: symbolic links are not allowed")
    _require(
        path.name == source_table_contract["filename"],
        f"{table_id}: filename differs from contract",
    )
    _require(
        source_table_contract["header_row"] == 1,
        f"{table_id}: only header row 1 is supported",
    )
    if source_table_contract["container"] == "CSV":
        return _read_csv_rows(path, source_table_contract)
    if source_table_contract["container"] == "XLSX":
        return _read_xlsx_rows(path, source_table_contract)
    raise SourceAdapterError(f"{table_id}: unsupported source container")


def extract_staging_tables_from_files(
    raw_root: Path,
    staging_contract: dict[str, Any],
    source_table_contract: dict[str, Any],
) -> dict[str, list[dict[str, Any]]]:
    """Read and extract only the tables declared as staging inputs."""

    _require(raw_root.is_dir(), "staging raw directory is missing")
    source_tables = {
        item["id"]: item for item in source_table_contract["table_contracts"]
    }
    workbook_tables: dict[str, list[dict[str, Any]]] = {}
    for source_table in source_tables.values():
        if source_table["container"] == "XLSX":
            workbook_tables.setdefault(source_table["filename"], []).append(
                source_table
            )
    for filename, contracts in workbook_tables.items():
        workbook_path = raw_root / filename
        _require(workbook_path.is_file(), f"{filename}: source workbook is missing")
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            expected_sheets = [contract["sheet_name"] for contract in contracts]
            _require(
                workbook.sheetnames == expected_sheets,
                f"{filename}: missing, unexpected, or reordered worksheet",
            )
        finally:
            workbook.close()
    extracted: dict[str, list[dict[str, Any]]] = {}
    for staging_table in staging_contract["staging_tables"]:
        source_id = staging_table["source_table_id"]
        _require(source_id in source_tables, f"unknown source table {source_id}")
        source_table = source_tables[source_id]
        source_rows = read_source_rows(
            raw_root / source_table["filename"], source_table
        )
        extracted[staging_table["id"]] = extract_staging_rows(
            source_rows, staging_table, source_table
        )
    return extracted
