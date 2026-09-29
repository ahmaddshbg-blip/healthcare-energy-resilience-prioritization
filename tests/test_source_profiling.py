from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from healthcare_resilience.contracts import (
    load_json,
    validate_observed_table_profile,
    validate_schema,
)
from healthcare_resilience.hashing import sha256_file
from healthcare_resilience.source_profiling import (
    SourceTableProfilingError,
    profile_source_tables,
    write_source_table_profile,
)

ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = ROOT / "configs"
FIXTURE_DIR = ROOT / "tests" / "fixtures"
SOURCE_CONTRACT_SHA256 = "2" * 64
TABLE_CONTRACT_SHA256 = "3" * 64


@pytest.fixture
def synthetic_raw_root(tmp_path: Path) -> Path:
    raw_root = tmp_path / "raw"
    raw_root.mkdir()
    (raw_root / "synthetic_measurements.csv").write_bytes(
        b"unit_id,amount\nU1,1.5\nU2,\nU3,3\n"
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Regions"
    sheet.append(["region_id", "label"])
    sheet.append(["R1", "North"])
    sheet.append(["R1", "South"])
    workbook.save(raw_root / "synthetic_regions.xlsx")
    workbook.close()
    return raw_root


def _contract() -> dict:
    return load_json(FIXTURE_DIR / "synthetic_source_table_contract.json")


def _profile(raw_root: Path) -> dict:
    return profile_source_tables(
        raw_root,
        _contract(),
        SOURCE_CONTRACT_SHA256,
        TABLE_CONTRACT_SHA256,
    )


def test_profiles_invented_csv_and_workbook(synthetic_raw_root: Path) -> None:
    profile = _profile(synthetic_raw_root)
    validate_schema(
        profile,
        load_json(CONFIG_DIR / "source_table_profile.schema.json"),
        "synthetic source-table profile",
    )
    validate_observed_table_profile(
        _contract(),
        profile,
        expected_source_contract_sha256=SOURCE_CONTRACT_SHA256,
        expected_source_table_contract_sha256=TABLE_CONTRACT_SHA256,
    )
    assert profile == load_json(
        FIXTURE_DIR / "synthetic_source_table_profile.json"
    )
    assert str(synthetic_raw_root) not in json.dumps(profile)


def test_profile_report_is_deterministic(synthetic_raw_root: Path) -> None:
    report_path = synthetic_raw_root.parent / "verification" / "profile.json"
    first = _profile(synthetic_raw_root)
    second = _profile(synthetic_raw_root)
    assert first == second
    write_source_table_profile(report_path, first)
    first_hash = sha256_file(report_path)
    write_source_table_profile(report_path, second)
    assert sha256_file(report_path) == first_hash
    assert not report_path.with_name(".profile.json.tmp").exists()


def test_invalid_utf8_stops_profiling(synthetic_raw_root: Path) -> None:
    csv_path = synthetic_raw_root / "synthetic_measurements.csv"
    csv_path.write_bytes(csv_path.read_bytes() + b"U4,\xff\n")
    with pytest.raises(SourceTableProfilingError, match="not valid UTF-8"):
        _profile(synthetic_raw_root)


def test_wrong_required_type_stops_profiling(synthetic_raw_root: Path) -> None:
    csv_path = synthetic_raw_root / "synthetic_measurements.csv"
    csv_path.write_text(
        "unit_id,amount\nU1,1.5\nU2,not-a-number\nU3,3\n",
        encoding="utf-8",
        newline="",
    )
    with pytest.raises(SourceTableProfilingError, match="cannot be parsed as NUMBER"):
        _profile(synthetic_raw_root)


def test_inconsistent_csv_width_stops_profiling(synthetic_raw_root: Path) -> None:
    csv_path = synthetic_raw_root / "synthetic_measurements.csv"
    csv_path.write_text(
        "unit_id,amount\nU1,1.5\nU2,\nU3,3,extra\n",
        encoding="utf-8",
        newline="",
    )
    with pytest.raises(SourceTableProfilingError, match="inconsistent data-row widths"):
        _profile(synthetic_raw_root)


def test_unexpected_worksheet_stops_profiling(synthetic_raw_root: Path) -> None:
    workbook_path = synthetic_raw_root / "synthetic_regions.xlsx"
    workbook = load_workbook(workbook_path)
    workbook.create_sheet("Unexpected")
    workbook.save(workbook_path)
    workbook.close()
    with pytest.raises(SourceTableProfilingError, match="unexpected, or reordered"):
        _profile(synthetic_raw_root)
