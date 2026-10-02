from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from healthcare_resilience.contracts import load_json
from healthcare_resilience.source_adapters import (
    SourceAdapterError,
    extract_staging_tables_from_files,
    read_source_rows,
)

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = ROOT / "tests" / "fixtures"


@pytest.fixture
def synthetic_raw_root(tmp_path: Path) -> Path:
    raw_root = tmp_path / "raw"
    raw_root.mkdir()
    (raw_root / "synthetic_measurements.csv").write_bytes(
        b"unit_id,amount\n01001,11\n 02020 ,\n03030,4.50\n"
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Regions"
    sheet.append(["region_id", "label"])
    sheet.append(["R-01", "Not Applicable"])
    sheet.append(["R-01", "Not Applicable"])
    workbook.save(raw_root / "synthetic_regions.xlsx")
    workbook.close()
    return raw_root


def _source_contract() -> dict:
    return load_json(FIXTURE_DIR / "synthetic_source_table_contract.json")


def _staging_contract() -> dict:
    return load_json(FIXTURE_DIR / "synthetic_staging_contract.json")


def _source_table(table_id: str) -> dict:
    return next(
        item for item in _source_contract()["table_contracts"]
        if item["id"] == table_id
    )


def test_csv_adapter_preserves_source_text_and_blank(
    synthetic_raw_root: Path,
) -> None:
    rows = read_source_rows(
        synthetic_raw_root / "synthetic_measurements.csv",
        _source_table("synthetic_measurements"),
    )
    assert rows == [
        {"unit_id": "01001", "amount": "11"},
        {"unit_id": " 02020 ", "amount": ""},
        {"unit_id": "03030", "amount": "4.50"},
    ]


def test_xlsx_adapter_preserves_repeated_rows(synthetic_raw_root: Path) -> None:
    rows = read_source_rows(
        synthetic_raw_root / "synthetic_regions.xlsx",
        _source_table("synthetic_regions"),
    )
    assert rows == [
        {"region_id": "R-01", "label": "Not Applicable"},
        {"region_id": "R-01", "label": "Not Applicable"},
    ]


def test_file_adapters_feed_source_preserving_extraction(
    synthetic_raw_root: Path,
) -> None:
    tables = extract_staging_tables_from_files(
        synthetic_raw_root,
        _staging_contract(),
        _source_contract(),
    )
    assert list(tables) == [
        "stg_synthetic_measurements",
        "stg_synthetic_regions",
    ]
    assert tables["stg_synthetic_measurements"][0] == {
        "source_row_number": 1,
        "unit_id": "01001",
        "amount": Decimal("11"),
    }
    assert tables["stg_synthetic_measurements"][1]["unit_id"] == " 02020 "
    assert tables["stg_synthetic_measurements"][1]["amount"] is None
    assert tables["stg_synthetic_regions"][0] == {
        "source_row_number": 1,
        "region_id": "R-01",
        "label": "Not Applicable",
    }


def test_adapter_rejects_ordered_header_drift(synthetic_raw_root: Path) -> None:
    path = synthetic_raw_root / "synthetic_measurements.csv"
    path.write_bytes(b"amount,unit_id\n11,01001\n,02020\n4.50,03030\n")
    with pytest.raises(SourceAdapterError, match="ordered header hash"):
        read_source_rows(path, _source_table("synthetic_measurements"))


def test_adapter_rejects_line_ending_drift(synthetic_raw_root: Path) -> None:
    path = synthetic_raw_root / "synthetic_measurements.csv"
    path.write_bytes(path.read_bytes().replace(b"\n", b"\r\n"))
    with pytest.raises(SourceAdapterError, match="line ending differs"):
        read_source_rows(path, _source_table("synthetic_measurements"))


def test_adapter_rejects_worksheet_row_count_drift(
    synthetic_raw_root: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Regions"
    sheet.append(["region_id", "label"])
    sheet.append(["R-01", "North"])
    workbook.save(synthetic_raw_root / "synthetic_regions.xlsx")
    workbook.close()
    with pytest.raises(SourceAdapterError, match="data row count"):
        read_source_rows(
            synthetic_raw_root / "synthetic_regions.xlsx",
            _source_table("synthetic_regions"),
        )


def test_file_extraction_rejects_unexpected_worksheet(
    synthetic_raw_root: Path,
) -> None:
    path = synthetic_raw_root / "synthetic_regions.xlsx"
    workbook = load_workbook(path)
    workbook.create_sheet("Unexpected")
    workbook.save(path)
    workbook.close()
    with pytest.raises(SourceAdapterError, match="unexpected, or reordered"):
        extract_staging_tables_from_files(
            synthetic_raw_root,
            _staging_contract(),
            _source_contract(),
        )
